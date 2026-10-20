from __future__ import annotations

import io
import json
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.models import Project
from app.storage import LocalStorage, set_storage

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
PDF = b"%PDF-1.7\n%test\n"


@pytest.fixture(autouse=True)
def tmp_storage(tmp_path):
    set_storage(LocalStorage(str(tmp_path)))
    yield tmp_path
    set_storage(None)


def _project(api, **kw):
    return api.post("/api/v1/projects", json={"name": "Water", "lifecycle": "testing", **kw}).json()


def test_manual_evidence_and_health_overview(api):
    p = _project(api)
    rec = api.post(
        "/api/v1/evidence",
        json={
            "project_id": p["id"],
            "kind": "calibration",
            "title": "Depth at 1.00 m",
            "result": "partial",
            "origin": "measured",
            "observed_at": "2026-10-19T10:00:00Z",
            "measured_value": "0.94",
            "reference_value": "1.00",
            "unit": "m",
        },
    )
    assert rec.status_code == 201, rec.text
    assert float(rec.json()["measured_value"]) == 0.94
    api.post(
        "/api/v1/evidence",
        json={
            "project_id": p["id"],
            "kind": "note",
            "title": "Told it works",
            "result": "info",
            "observed_at": "2026-10-19T11:00:00Z",
        },
    )
    overview = {o["name"]: o for o in api.get("/api/v1/health-overview").json()}["Water"]
    assert overview["verified_count"] == 1 and overview["reported_count"] == 1
    assert overview["last_title"] == "Depth at 1.00 m" and overview["stale"] is False
    bad = api.post(
        "/api/v1/evidence",
        json={
            "project_id": p["id"],
            "kind": "ci_run",
            "title": "x",
            "result": "pass",
            "origin": "automated",
            "observed_at": "2026-10-19T10:00:00Z",
        },
    )
    assert bad.status_code == 422  # manual records can't claim to be automated
    future = api.post(
        "/api/v1/evidence",
        json={
            "project_id": p["id"],
            "kind": "note",
            "title": "x",
            "result": "info",
            "observed_at": "2099-01-01T00:00:00Z",
        },
    )
    assert future.status_code == 422
    assert api.get(f"/api/v1/evidence?project_id={p['id']}&origin=measured").json()["total"] == 1


def test_document_upload_versions_and_download(api):
    p = _project(api)
    res = api.post(
        "/api/v1/documents/upload",
        files={"file": ("wiring.png", PNG, "image/png")},
        data={"project_id": p["id"], "category": "wiring"},
    )
    assert res.status_code == 201, res.text
    doc = res.json()
    assert doc["title"] == "wiring.png" and doc["versions"][0]["content_type"] == "image/png"
    v2 = api.post(f"/api/v1/documents/{doc['id']}/versions", files={"file": ("../../etc/passwd.pdf", PDF, "x")})
    assert v2.status_code == 201 and v2.json()["version_no"] == 2 and v2.json()["filename"] == "passwd.pdf"
    dl = api.get(f"/api/v1/documents/{doc['id']}/versions/{v2.json()['id']}/download")
    assert dl.content == PDF and dl.headers["content-disposition"].startswith("attachment")
    assert "sandbox" in dl.headers["content-security-policy"]
    inline = api.get(f"/api/v1/documents/{doc['id']}/versions/{v2.json()['id']}/download?inline=true")
    assert inline.headers["content-disposition"].startswith("inline")


def test_rejects_html_svg_and_mislabelled_files(api, tmp_storage):
    p = _project(api)
    for name, data in [
        ("page.html", b"<html><script>alert(1)</script>"),
        ("x.svg", b"<svg onload=alert(1)>"),
        ("fake.png", b"MZ\x90\x00 not an image"),
        ("t.txt", b"bad\x00binary"),
    ]:
        res = api.post(
            "/api/v1/documents/upload", files={"file": (name, data, "image/png")}, data={"project_id": p["id"]}
        )
        assert res.status_code == 422, name
    assert api.get("/api/v1/documents").json()["total"] == 0
    assert list(tmp_storage.iterdir()) == []


def test_office_and_text_detection(api):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", "<w/>")
    res = api.post("/api/v1/documents/upload", files={"file": ("srs.docx", buf.getvalue(), "x")})
    assert res.json()["versions"][0]["content_type"].endswith("wordprocessingml.document")
    md = api.post("/api/v1/documents/upload", files={"file": ("notes.md", "# Hi ✓".encode(), "x")})
    assert md.json()["versions"][0]["content_type"] == "text/markdown"


def test_document_size_limit(api, monkeypatch):
    from app.core.config import get_settings

    monkeypatch.setattr(get_settings(), "max_upload_mb", 1)
    big = PDF + b"0" * (1024 * 1024 + 10)
    assert api.post("/api/v1/documents/upload", files={"file": ("big.pdf", big, "x")}).status_code == 422


def test_knowledge_base_search_and_delete(api, tmp_storage):
    p = _project(api)
    note = api.post(
        "/api/v1/documents",
        json={
            "project_id": p["id"],
            "title": "LoRa decision",
            "category": "decision",
            "body_markdown": "We chose **SX1278** at 433 MHz.",
        },
    ).json()
    api.post(f"/api/v1/documents/{note['id']}/versions", files={"file": ("d.pdf", PDF, "x")})
    assert api.get("/api/v1/documents?q=sx1278").json()["total"] == 1
    assert api.get("/api/v1/documents?q=nothing").json()["total"] == 0
    assert len(list(tmp_storage.iterdir())) == 1
    assert api.delete(f"/api/v1/documents/{note['id']}").status_code == 204
    assert list(tmp_storage.iterdir()) == []  # stored file removed too


def test_business_entries_and_summary_keep_currencies_apart(api):
    p = _project(api, name="E-Moto")
    for kind, amount, cur, status in [
        ("cost", "2500000.00", "TZS", "estimated"),
        ("cost", "1500000", "TZS", "estimated"),
        ("funding_need", "12000", "USD", "open"),
    ]:
        res = api.post(
            "/api/v1/business",
            json={
                "project_id": p["id"],
                "kind": kind,
                "title": "x",
                "amount": amount,
                "currency": cur,
                "status": status,
            },
        )
        assert res.status_code == 201, res.text
    assert (
        api.post(
            "/api/v1/business", json={"project_id": p["id"], "kind": "cost", "title": "x", "amount": "-1"}
        ).status_code
        == 422
    )
    assert (
        api.post(
            "/api/v1/business", json={"project_id": p["id"], "kind": "cost", "title": "x", "currency": "usd"}
        ).status_code
        == 422
    )
    summary = {(s["currency"], s["kind"]): s for s in api.get("/api/v1/business/summary").json()}
    assert summary[("TZS", "cost")]["total"] == "4000000.00" and summary[("TZS", "cost")]["count"] == 2
    assert summary[("USD", "funding_need")]["total"] == "12000.00"


def test_export_json_has_data_but_no_secrets(api, db):
    from app.models import User
    from app.services.auth import store_github_token

    _project(api)
    store_github_token(db, db.scalar(select(User)), "gho_supersecret", "repo")
    db.commit()
    res = api.get("/api/v1/export/json")
    body = res.content.decode()
    assert "gho_supersecret" not in body and "access_token_enc" not in body and "token_hash" not in body
    data = json.loads(body)
    assert data["tables"]["projects"][0]["name"] == "Water"


def test_export_register_round_trips_through_import(api):
    from factories import make_register

    content = make_register()
    sha = api.post("/api/v1/import/register/preview", files={"file": ("r.xlsx", content)}).json()["sha256"]
    api.post("/api/v1/import/register/commit", files={"file": ("r.xlsx", content)}, data={"sha256": sha})
    exported = api.get("/api/v1/export/register.xlsx").content
    wb = load_workbook(io.BytesIO(exported))
    assert wb.sheetnames == ["Master Register", "Priority Plan"]
    preview = api.post("/api/v1/import/register/preview", files={"file": ("e.xlsx", exported)}).json()
    assert preview["summary"]["unchanged"] == 7, [r for r in preview["rows"] if r["action"] != "unchanged"]


def test_delete_everything_requires_exact_confirmation(api, db, tmp_storage):
    _project(api)
    api.post("/api/v1/documents/upload", files={"file": ("a.pdf", PDF, "x")})
    assert api.post("/api/v1/account/delete-all", json={"confirm": "yes"}).status_code == 422
    res = api.post("/api/v1/account/delete-all", json={"confirm": "DELETE EVERYTHING"})
    assert res.status_code == 200
    assert db.scalar(select(func.count()).select_from(Project)) == 0
    assert list(tmp_storage.iterdir()) == []
    assert api.get("/api/v1/auth/me").status_code == 401
