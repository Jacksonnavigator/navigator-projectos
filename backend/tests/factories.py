"""Builds small .xlsx registers shaped like the owner's real 'Master Project Register'."""

from __future__ import annotations

import io

from openpyxl import Workbook

MASTER_HEADER = [
    "ID",
    "Project",
    "Category",
    "Type",
    "Status",
    "Status confidence",
    "Purpose / scope",
    "Known tech stack",
    "Repository / URL",
    "Known remaining work",
    "Business potential",
    "Suggested next action",
    "Evidence / notes",
]


def master_row(rid, name, status="Active development", repo="Not confirmed", notes="", next_action="Do the thing"):
    return [
        rid,
        name,
        "Software",
        "Product",
        status,
        "Prior report; unverified",
        f"Purpose of {name}",
        "FastAPI",
        repo,
        "Remaining work",
        "High",
        next_action,
        notes,
    ]


DEFAULT_ROWS = [
    master_row(
        1, "DevGuardian AI", "Active development", "Local path previously referenced: C:/Users/x/Desktop/DevGuardian"
    ),
    master_row(3, "Self-Hosted AI Infrastructure / Local LLM Server", "Exploration / setup"),
    master_row(
        4,
        "Tanzania Bird Species Identifier",
        "Prototype / model troubleshooting",
        "Hugging Face: https://huggingface.co/Jacksonnavigator7/Birds_Classifications",
    ),
    master_row(
        17,
        "Meza Order",
        "Deployed; QA and security follow-up needed",
        "https://meza-order.onrender.com/ ; https://github.com/Jacksonnavigator/restaurant",
    ),
    master_row(
        24,
        "AI-Powered Wildlife / Bird Monitoring Device",
        "Concept / potential extension",
        notes="Overlaps with the Tanzania Bird Species Identifier.",
    ),
    master_row(
        31,
        "Local AI API / Developer Infrastructure",
        "Part of AI server direction",
        notes="Avoid duplicate counting with project 3.",
        next_action="Treat as a component of the self-hosted AI infrastructure unless it becomes separate.",
    ),
    master_row(
        33,
        "Application Security / CI-CD Reliability Tooling",
        "Partly implemented in DevGuardian",
        notes="Not necessarily a separate product.",
        next_action="Merge its backlog into DevGuardian.",
    ),
]

PRIORITY_HEADER = ["Rank", "Project", "Suggested priority", "Reason", "First next action"]
DEFAULT_PRIORITIES = [
    [1, "DevGuardian AI", "P1 — Immediate", "Existing implementation.", "Run tests."],
    [2, "Meza Order", "P1 — Immediate", "Deployed application.", "Fix security findings."],
    [7, "Bird Species Identifier", "P2 — Next", "Distinct AI project.", "Fix inference."],
    [10, "RestaurantOS", "P3 — Decide", "May overlap with Meza Order.", "Decide."],
]


def make_register(rows=None, priorities=None, header=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Register"
    ws.append(header or MASTER_HEADER)
    for row in rows if rows is not None else DEFAULT_ROWS:
        ws.append(row)
    plan = wb.create_sheet("Priority Plan")
    plan.append(PRIORITY_HEADER)
    for row in priorities if priorities is not None else DEFAULT_PRIORITIES:
        plan.append(row)
    guide = wb.create_sheet("How to Use")
    guide.append(["Topic", "Guidance"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
