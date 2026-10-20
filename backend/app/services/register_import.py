"""Import the Excel project register (Master Register + Priority Plan sheets).

Flow: preview (no writes) -> user reviews -> commit with the same file (checked by SHA-256).
The server keeps no state between the two calls, because the free Render instance may sleep in between.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import ValidationFailed
from app.models import Project, ProjectLink, ProjectRelation
from app.services.projects import extract_links, record_activity

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000

# Canonical column -> accepted header spellings (normalised: lowercase, alphanumerics only).
MASTER_COLUMNS: dict[str, tuple[str, ...]] = {
    "register_id": ("id",),
    "name": ("project", "projectname", "name"),
    "category": ("category",),
    "project_type": ("type", "projecttype"),
    "reported_status": ("status",),
    "status_confidence": ("statusconfidence",),
    "purpose": ("purposescope", "purpose", "scope"),
    "tech_stack": ("knowntechstack", "techstack"),
    "repository_text": ("repositoryurl", "repository", "url"),
    "remaining_work": ("knownremainingwork", "remainingwork"),
    "business_potential": ("businesspotential",),
    "next_action": ("suggestednextaction", "nextaction"),
    "evidence_notes": ("evidencenotes", "evidence", "notes"),
}
REQUIRED = ("name",)
PRIORITY_COLUMNS: dict[str, tuple[str, ...]] = {
    "rank": ("rank",),
    "name": ("project",),
    "priority": ("suggestedpriority", "priority"),
    "reason": ("reason",),
    "first_action": ("firstnextaction", "nextaction"),
}
# Fields the import owns; on re-import these are overwritten from the spreadsheet.
IMPORTED_FIELDS = tuple(k for k in MASTER_COLUMNS if k != "register_id")
_CAMEL_ALIAS = re.compile(r"^[A-Z][a-z0-9]+[A-Z0-9][A-Za-z0-9]*$")  # DevGuardian, RestaurantOS, My3DWorld
_PROJECT_REF = re.compile(r"\bproject\s+#?(\d{1,4})\b", re.IGNORECASE)
_PARENT_HINT = re.compile(
    r"component of|part of|belongs under|merge (?:its|the) backlog into|track(?:ed)? as .{0,40}modules|subproject",
    re.IGNORECASE,
)


def _norm(header: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def infer_lifecycle(status: str | None) -> str:
    """Best-effort mapping of the free-text register status to a lifecycle stage (editable afterwards)."""
    s = (status or "").lower()
    if not s:
        return "planning"
    if "deployed" in s:
        return "deployed"
    if any(w in s for w in ("unknown", "unclear", "historical")):
        return "paused"
    if any(w in s for w in ("testing", "validation", "integration", "troubleshooting", "configuration")):
        return "testing"
    if any(w in s for w in ("requirements", "planning", "blueprint", "concept", "design", "proposal")):
        return "planning"
    if any(w in s for w in ("exploration", "idea", "research notes")):
        return "idea"
    if any(w in s for w in ("active", "development", "prototype", "ongoing", "complete", "installed", "learning")):
        return "active"
    return "planning"


def parse_priority(text: str | None) -> tuple[str | None, str | None]:
    """'P1 — Immediate' -> ('P1', 'Immediate')."""
    if not text:
        return None, None
    m = re.match(r"\s*(P[1-4])\s*[—–\-:]*\s*(.*)$", text)
    if not m:
        return None, text.strip()[:100]
    return m.group(1), (m.group(2).strip() or None)


@dataclass
class ParsedRow:
    row_number: int
    values: dict[str, Any]
    problems: list[str] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    sha256: str
    rows: list[ParsedRow]
    priority_rows: list[dict[str, Any]]
    warnings: list[str]


def _header_map(header_row: tuple[Any, ...], columns: dict[str, tuple[str, ...]]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalised = [_norm(h) for h in header_row]
    for canonical, spellings in columns.items():
        for spelling in spellings:
            if spelling in normalised:
                mapping[canonical] = normalised.index(spelling)
                break
    return mapping


def _find_sheet(wb: Any, *names: str) -> Any:
    wanted = {_norm(n) for n in names}
    for ws in wb.worksheets:
        if _norm(ws.title) in wanted:
            return ws
    return None


def parse_workbook(content: bytes) -> ParsedWorkbook:
    if len(content) > MAX_FILE_BYTES:
        raise ValidationFailed("The file is larger than 5 MB")
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ValidationFailed("This is not an .xlsx file. Save the register as Excel Workbook (.xlsx) and retry.")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is a user error
        raise ValidationFailed("The workbook could not be read. Is it a valid .xlsx file?") from exc
    warnings: list[str] = []
    master = _find_sheet(wb, "Master Register", "Register", "Projects") or wb.worksheets[0]
    rows_iter = master.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        raise ValidationFailed(f"Sheet '{master.title}' is empty")
    mapping = _header_map(header, MASTER_COLUMNS)
    missing = [c for c in REQUIRED if c not in mapping]
    if missing:
        raise ValidationFailed(
            f"Sheet '{master.title}' has no 'Project' column. Found columns: " + ", ".join(str(h) for h in header if h)
        )
    unknown = [str(h) for i, h in enumerate(header) if h and i not in mapping.values()]
    if unknown:
        warnings.append("Ignored columns: " + ", ".join(unknown))

    rows: list[ParsedRow] = []
    for idx, raw in enumerate(rows_iter, start=2):
        if idx - 1 > MAX_ROWS:
            raise ValidationFailed(f"More than {MAX_ROWS} rows; split the register")
        values = {key: _clean(raw[col]) if col < len(raw) else None for key, col in mapping.items()}
        if all(v is None for v in values.values()):
            continue
        row = ParsedRow(row_number=idx, values=values)
        name = values.get("name")
        if not name:
            row.problems.append("Missing project name")
        elif len(str(name)) > 200:
            row.problems.append("Project name is longer than 200 characters")
        rid = values.get("register_id")
        if rid is not None and not isinstance(rid, int):
            if str(rid).isdigit():
                values["register_id"] = int(rid)
            else:
                row.problems.append(f"ID '{rid}' is not a whole number; it will be ignored")
                values["register_id"] = None
        for key, value in list(values.items()):
            if key not in ("register_id",) and value is not None and not isinstance(value, str):
                values[key] = str(value)
        rows.append(row)

    # Duplicate checks inside the file.
    seen_ids: dict[int, int] = {}
    seen_names: dict[str, int] = {}
    for row in rows:
        rid = row.values.get("register_id")
        if rid is not None:
            if rid in seen_ids:
                row.problems.append(f"Duplicate ID {rid} (also on row {seen_ids[rid]})")
            seen_ids.setdefault(rid, row.row_number)
        name = (row.values.get("name") or "").lower()
        if name:
            if name in seen_names:
                row.problems.append(f"Duplicate project name (also on row {seen_names[name]})")
            seen_names.setdefault(name, row.row_number)

    priority_rows: list[dict[str, Any]] = []
    plan = _find_sheet(wb, "Priority Plan", "Priorities")
    if plan is not None:
        plan_iter = plan.iter_rows(values_only=True)
        plan_header = next(plan_iter, None) or ()
        pmap = _header_map(plan_header, PRIORITY_COLUMNS)
        if "name" in pmap:
            for raw in plan_iter:
                vals = {k: _clean(raw[c]) if c < len(raw) else None for k, c in pmap.items()}
                if vals.get("name"):
                    priority_rows.append(vals)
        else:
            warnings.append(f"Sheet '{plan.title}' has no 'Project' column; priorities were not imported")
    wb.close()
    return ParsedWorkbook(hashlib.sha256(content).hexdigest(), rows, priority_rows, warnings)


def _simplify(name: str) -> str:
    return re.sub(r"[^a-z0-9 ]", " ", name.lower()).split("(")[0].strip()


def match_priority_rows(rows: list[ParsedRow], priority_rows: list[dict[str, Any]]) -> tuple[dict[int, dict], list]:
    """Map Priority Plan entries to register rows by exact, then unambiguous partial, name match."""
    by_name = {(r.values.get("name") or "").lower(): r for r in rows if r.values.get("name")}
    matched: dict[int, dict[str, Any]] = {}
    unmatched: list[dict[str, Any]] = []
    for prow in priority_rows:
        pname = str(prow["name"]).strip()
        row = by_name.get(pname.lower())
        how = "exact"
        if row is None:
            simple = " ".join(_simplify(pname).split())
            candidates = [
                r
                for r in rows
                if r.values.get("name")
                and (
                    simple in " ".join(_simplify(r.values["name"]).split())
                    or " ".join(_simplify(r.values["name"]).split()) in simple
                )
            ]
            row = candidates[0] if len(candidates) == 1 else None
            how = "partial"
        if row is None:
            unmatched.append({"name": pname, "reason": "No single project with a matching name"})
            continue
        if row.row_number in matched:
            unmatched.append({"name": pname, "reason": f"Row {row.row_number} already has a priority"})
            continue
        matched[row.row_number] = {**prow, "match": how}
    return matched, unmatched


def _aliases(name: str) -> list[str]:
    out = {name.lower()}
    for sep in (" / ", " (", "/"):
        if sep in name:
            out.add(name.split(sep)[0].strip().lower())
    first = name.split()[0] if name.split() else ""
    if _CAMEL_ALIAS.match(first):
        out.add(first.lower())
    return [a for a in out if len(a) >= 5]


def relation_suggestions(rows: list[ParsedRow]) -> list[dict[str, Any]]:
    """Suggest parent/related links from notes like 'part of DevGuardian' or 'duplicate counting with project 3'."""
    by_id = {r.values["register_id"]: r for r in rows if r.values.get("register_id") is not None}
    alias_index = [(alias, r) for r in rows if r.values.get("name") for alias in _aliases(r.values["name"])]
    suggestions: dict[tuple[int, int], dict[str, Any]] = {}
    for row in rows:
        rid = row.values.get("register_id")
        if rid is None:
            continue
        text = " ".join(
            str(row.values.get(k) or "") for k in ("evidence_notes", "next_action", "remaining_work", "repository_text")
        )
        if not text.strip():
            continue
        lowered = text.lower()
        targets: set[int] = set()
        for m in _PROJECT_REF.finditer(text):
            ref = int(m.group(1))
            if ref in by_id and ref != rid:
                targets.add(ref)
        for alias, other in alias_index:
            oid = other.values.get("register_id")
            if oid is None or oid == rid:
                continue
            if re.search(r"\b" + re.escape(alias) + r"\b", lowered):
                targets.add(oid)
        parent_hint = bool(_PARENT_HINT.search(text))
        for target in targets:
            kind = "parent" if parent_hint else "related"
            key = (rid, target) if kind == "parent" else tuple(sorted((rid, target)))
            if key in suggestions:
                continue
            suggestions[key] = {
                "kind": kind,
                "child_register_id" if kind == "parent" else "a_register_id": rid,
                "parent_register_id" if kind == "parent" else "b_register_id": target,
                "description": (
                    f"#{rid} {row.values['name']} → part of #{target} {by_id[target].values['name']}"
                    if kind == "parent"
                    else f"#{rid} {row.values['name']} ↔ related to #{target} {by_id[target].values['name']}"
                ),
            }
    return list(suggestions.values())


def _existing_for(db: Session, values: dict[str, Any]) -> Project | None:
    rid = values.get("register_id")
    if rid is not None:
        project = db.scalar(select(Project).where(Project.register_id == rid))
        if project:
            return project
    name = values.get("name")
    if name:
        return db.scalar(select(Project).where(func.lower(Project.name) == name.lower()))
    return None


def _target_values(row: ParsedRow, priority: dict[str, Any] | None) -> dict[str, Any]:
    vals = {k: row.values.get(k) for k in IMPORTED_FIELDS}
    if priority:
        code, label = parse_priority(priority.get("priority"))
        vals["priority"] = code
        vals["priority_label"] = label
        rank = priority.get("rank")
        vals["priority_rank"] = int(str(rank)) if isinstance(rank, int) or str(rank or "").isdigit() else None
        vals["priority_reason"] = priority.get("reason")
        if priority.get("first_action") and not vals.get("next_action"):
            vals["next_action"] = priority["first_action"]
    return vals


def build_preview(db: Session, parsed: ParsedWorkbook) -> dict[str, Any]:
    matched, unmatched = match_priority_rows(parsed.rows, parsed.priority_rows)
    rows_out = []
    counts = {"create": 0, "update": 0, "unchanged": 0, "skip": 0}
    for row in parsed.rows:
        blocking = [p for p in row.problems if "ignored" not in p]
        priority = matched.get(row.row_number)
        target = _target_values(row, priority)
        existing = None if blocking else _existing_for(db, row.values)
        if blocking:
            action, changes = "skip", list[dict[str, Any]]()
        elif existing is None:
            action = "create"
            changes = list[dict[str, Any]]()
        else:
            changes = [
                {"field": k, "old": getattr(existing, k), "new": v}
                for k, v in target.items()
                if (getattr(existing, k) or None) != (v or None)
            ]
            if row.values.get("register_id") is not None and existing.register_id != row.values["register_id"]:
                changes.append({"field": "register_id", "old": existing.register_id, "new": row.values["register_id"]})
            action = "update" if changes else "unchanged"
        counts[action] += 1
        rows_out.append(
            {
                "row_number": row.row_number,
                "register_id": row.values.get("register_id"),
                "name": row.values.get("name"),
                "action": action,
                "lifecycle": infer_lifecycle(row.values.get("reported_status")) if action == "create" else None,
                "priority": target.get("priority"),
                "priority_match": priority.get("match") if priority else None,
                "links": [{"kind": k, "url": u} for k, u in extract_links(row.values.get("repository_text"))],
                "changes": [{**c, "old": _jsonable(c["old"]), "new": _jsonable(c["new"])} for c in changes],
                "problems": row.problems,
            }
        )
    return {
        "sha256": parsed.sha256,
        "summary": counts,
        "rows": rows_out,
        "priority_unmatched": unmatched,
        "suggestions": relation_suggestions(parsed.rows),
        "warnings": parsed.warnings,
    }


def _jsonable(value: Any) -> Any:
    return value if value is None or isinstance(value, (str, int, float, bool)) else str(value)


def commit_import(
    db: Session, parsed: ParsedWorkbook, expected_sha256: str, apply_suggestions: list[dict[str, Any]]
) -> dict[str, Any]:
    if parsed.sha256 != expected_sha256:
        raise ValidationFailed("The file changed since the preview. Preview it again before importing.")
    matched, _ = match_priority_rows(parsed.rows, parsed.priority_rows)
    created = updated = skipped = 0
    by_register: dict[int, Project] = {}
    for row in parsed.rows:
        if [p for p in row.problems if "ignored" not in p]:
            skipped += 1
            continue
        target = _target_values(row, matched.get(row.row_number))
        project = _existing_for(db, row.values)
        if project is None:
            project = Project(
                register_id=row.values.get("register_id"),
                lifecycle=infer_lifecycle(row.values.get("reported_status")),
                **target,
            )
            db.add(project)
            db.flush()
            created += 1
            record_activity(
                db, "project.imported", f"Imported '{project.name}' from the register", project_id=project.id
            )
        else:
            changed = False
            for key, value in target.items():
                if (getattr(project, key) or None) != (value or None):
                    setattr(project, key, value)
                    changed = True
            rid = row.values.get("register_id")
            if rid is not None and project.register_id != rid:
                project.register_id = rid
                changed = True
            updated += int(changed)
        existing_urls = set(db.scalars(select(ProjectLink.url).where(ProjectLink.project_id == project.id)))
        for kind, url in extract_links(row.values.get("repository_text")):
            if url not in existing_urls:
                db.add(ProjectLink(project_id=project.id, kind=kind, url=url))
                existing_urls.add(url)
        if project.register_id is not None:
            by_register[project.register_id] = project
    db.flush()

    applied = 0
    for s in apply_suggestions:
        if s.get("kind") == "parent":
            child = by_register.get(s.get("child_register_id"))  # type: ignore[arg-type]
            parent = by_register.get(s.get("parent_register_id"))  # type: ignore[arg-type]
            if child and parent and child.id != parent.id and parent.parent_id != child.id:
                child.parent_id = parent.id
                applied += 1
        elif s.get("kind") == "related":
            a = by_register.get(s.get("a_register_id"))  # type: ignore[arg-type]
            b = by_register.get(s.get("b_register_id"))  # type: ignore[arg-type]
            if a and b and a.id != b.id:
                x, y = sorted([a.id, b.id])
                exists = db.scalar(
                    select(ProjectRelation.id).where(
                        ProjectRelation.project_a_id == x, ProjectRelation.project_b_id == y
                    )
                )
                if not exists:
                    db.add(ProjectRelation(project_a_id=x, project_b_id=y, note="Suggested by register import"))
                applied += 1
    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "links_applied": applied}
